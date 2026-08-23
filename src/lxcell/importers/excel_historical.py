"""Dry-run importer for historical yearly Excel workbooks."""

from __future__ import annotations

import hashlib
import unicodedata
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from lxcell.enums.core_enums import Direction

DATE_HEADERS = {"fecha", "date", "dia", "día"}
IGNORED_HEADERS = {
    "total",
    "totales",
    "comentario",
    "comentarios",
    "comment",
    "comments",
    "nota",
    "notas",
    "observaciones",
}
TRACKING_SHEET_NAME = "Seguimiento"
TRACKING_TOLERANCE_MINOR = 1
TRACKING_MONTH_HEADERS = {"mes", "month", "fecha", "date"}
TRACKING_CATEGORY_HEADERS = {"categoria", "categoría", "category"}
MONTH_NAMES = {
    "enero": 1,
    "ene": 1,
    "january": 1,
    "jan": 1,
    "febrero": 2,
    "feb": 2,
    "february": 2,
    "marzo": 3,
    "mar": 3,
    "march": 3,
    "abril": 4,
    "abr": 4,
    "april": 4,
    "apr": 4,
    "mayo": 5,
    "may": 5,
    "junio": 6,
    "jun": 6,
    "june": 6,
    "julio": 7,
    "jul": 7,
    "july": 7,
    "agosto": 8,
    "ago": 8,
    "august": 8,
    "aug": 8,
    "septiembre": 9,
    "setiembre": 9,
    "sep": 9,
    "sept": 9,
    "september": 9,
    "octubre": 10,
    "oct": 10,
    "october": 10,
    "noviembre": 11,
    "nov": 11,
    "november": 11,
    "diciembre": 12,
    "dic": 12,
    "december": 12,
    "dec": 12,
}


@dataclass(frozen=True)
class HistoricalExcelTransactionCandidate:
    """One transaction candidate read from a historical Excel workbook."""

    row_number_source: int
    column_name_source: str
    transaction_date: date
    source_category_name: str
    amount_minor: int
    direction: Direction
    amount_raw: str
    description_raw: str | None
    payload_raw: dict[str, str | None]
    source_amount_minor: int | None = None
    source_amount_decimal: Decimal | None = None
    source_column_kind: str = "expense"

    @property
    def signed_amount_minor(self) -> int:
        if self.direction == Direction.INFLOW:
            return self.amount_minor
        if self.direction == Direction.OUTFLOW:
            return -self.amount_minor
        return 0

    @property
    def source_signed_amount_minor(self) -> int:
        if self.source_amount_minor is not None:
            return self.source_amount_minor
        return self.signed_amount_minor


@dataclass(frozen=True)
class HistoricalExcelPreview:
    """Read-only preview of a historical Excel import."""

    source_file_name: str
    source_file_hash: str
    sheet_name: str
    header_row_number: int
    date_column_name: str
    candidates: tuple[HistoricalExcelTransactionCandidate, ...]
    ignored_row_numbers: tuple[int, ...]
    tracking_validation: HistoricalExcelTrackingValidation | None = None

    @property
    def transaction_count(self) -> int:
        return len(self.candidates)

    @property
    def source_categories(self) -> tuple[str, ...]:
        return tuple(sorted({candidate.source_category_name for candidate in self.candidates}))

    @property
    def totals_by_category_minor(self) -> dict[str, int]:
        totals: dict[str, int] = defaultdict(int)
        for candidate in self.candidates:
            totals[candidate.source_category_name] += candidate.signed_amount_minor
        return dict(sorted(totals.items()))

    @property
    def totals_by_month_minor(self) -> dict[str, int]:
        totals: dict[str, int] = defaultdict(int)
        for candidate in self.candidates:
            month_key = candidate.transaction_date.strftime("%Y-%m")
            totals[month_key] += candidate.signed_amount_minor
        return dict(sorted(totals.items()))

    @property
    def source_totals_by_category_minor(self) -> dict[str, int]:
        totals: dict[str, Decimal] = defaultdict(Decimal)
        for candidate in self.candidates:
            totals[candidate.source_category_name] += candidate_source_amount_decimal(
                candidate
            )
        return {
            category_name: decimal_to_minor_units(amount)
            for category_name, amount in sorted(totals.items())
        }

    @property
    def source_totals_by_month_minor(self) -> dict[str, int]:
        totals: dict[str, Decimal] = defaultdict(Decimal)
        for candidate in self.candidates:
            month_key = candidate.transaction_date.strftime("%Y-%m")
            totals[month_key] += candidate_source_amount_decimal(candidate)
        return {
            month_key: decimal_to_minor_units(amount)
            for month_key, amount in sorted(totals.items())
        }


@dataclass(frozen=True)
class HistoricalExcelTrackingComparison:
    """One aggregate comparison between Registro and Seguimiento."""

    month_key: str
    source_category_name: str
    registro_amount_minor: int
    seguimiento_amount_minor: int

    @property
    def difference_minor(self) -> int:
        return self.registro_amount_minor - self.seguimiento_amount_minor

    @property
    def status(self) -> str:
        if abs(self.difference_minor) <= TRACKING_TOLERANCE_MINOR:
            return "ok"
        return "difference"


@dataclass(frozen=True)
class HistoricalExcelTrackingValidation:
    """Read-only validation of Registro aggregates against Seguimiento."""

    sheet_name: str
    comparisons: tuple[HistoricalExcelTrackingComparison, ...]
    registro_only_categories: tuple[str, ...]
    seguimiento_only_categories: tuple[str, ...]

    @property
    def matched_count(self) -> int:
        return len(self.comparisons)

    @property
    def difference_count(self) -> int:
        return sum(1 for comparison in self.comparisons if comparison.status != "ok")

    @property
    def ok_count(self) -> int:
        return self.matched_count - self.difference_count


@dataclass(frozen=True)
class TrackingTotalsParseResult:
    """Parsed tracking totals and category headers from Seguimiento."""

    totals: dict[tuple[str, str], int]
    categories: tuple[str, ...]


class HistoricalExcelDryRunImporter:
    """Parse historical Excel workbooks without writing to the database."""

    def __init__(
        self,
        *,
        sheet_name: str = "Registro",
        tracking_sheet_name: str = TRACKING_SHEET_NAME,
    ) -> None:
        self.sheet_name = sheet_name
        self.tracking_sheet_name = tracking_sheet_name

    def preview(self, workbook_path: str | Path) -> HistoricalExcelPreview:
        path = Path(workbook_path)
        workbook_hash = file_sha256(path)
        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=True,
        )
        try:
            if self.sheet_name not in workbook.sheetnames:
                raise ValueError(f"Workbook does not contain sheet '{self.sheet_name}'.")
            sheet = workbook[self.sheet_name]
            header_row_number, headers, header_cells = find_header_row(sheet.iter_rows())
            date_column_index = find_date_column_index(headers)
            comment_column_index = find_comment_column_index(headers)
            category_columns = category_column_indexes(headers, date_column_index)
            source_column_kinds = {
                column_index: infer_source_column_kind(
                    headers[column_index],
                    header_cells[column_index],
                )
                for column_index in category_columns
            }

            candidates: list[HistoricalExcelTransactionCandidate] = []
            ignored_row_numbers: list[int] = []
            for row_number, row in enumerate(
                sheet.iter_rows(
                    min_row=header_row_number + 1,
                ),
                start=header_row_number + 1,
            ):
                transaction_date = parse_excel_date(cell_value(row, date_column_index))
                if transaction_date is None:
                    if any(cell_value(row, index) is not None for index in range(len(row))):
                        ignored_row_numbers.append(row_number)
                    continue

                description_raw = string_or_none(cell_value(row, comment_column_index))
                row_created_candidate = False
                for column_index in category_columns:
                    amount = parse_excel_amount(cell_value(row, column_index))
                    if amount is None or amount == 0:
                        continue
                    category_name = str(headers[column_index]).strip()
                    source_column_kind = source_column_kinds[column_index]
                    candidates.append(
                        HistoricalExcelTransactionCandidate(
                            row_number_source=row_number,
                            column_name_source=category_name,
                            transaction_date=transaction_date,
                            source_category_name=category_name,
                            amount_minor=decimal_to_minor_units(abs(amount)),
                            direction=direction_from_source_amount(
                                amount,
                                source_column_kind=source_column_kind,
                            ),
                            amount_raw=str(cell_value(row, column_index)),
                            description_raw=description_raw,
                            payload_raw=payload_from_row(headers, row),
                            source_amount_minor=decimal_to_minor_units(amount),
                            source_amount_decimal=amount,
                            source_column_kind=source_column_kind,
                        )
                    )
                    row_created_candidate = True
                if not row_created_candidate:
                    ignored_row_numbers.append(row_number)

            return HistoricalExcelPreview(
                source_file_name=path.name,
                source_file_hash=workbook_hash,
                sheet_name=self.sheet_name,
                header_row_number=header_row_number,
                date_column_name=str(headers[date_column_index]),
                candidates=tuple(candidates),
                ignored_row_numbers=tuple(ignored_row_numbers),
                tracking_validation=build_tracking_validation(
                    workbook,
                    tracking_sheet_name=self.tracking_sheet_name,
                    candidates=tuple(candidates),
                ),
            )
        finally:
            workbook.close()


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def find_header_row(rows) -> tuple[int, tuple[str, ...], tuple[Any, ...]]:
    for row_number, row in enumerate(rows, start=1):
        headers = tuple(
            "" if cell_value(row, index) is None else str(cell_value(row, index)).strip()
            for index in range(len(row))
        )
        if find_date_column_index(headers) is not None and any(
            header for header in headers if normalize_header(header) not in DATE_HEADERS
        ):
            return row_number, headers, row
    raise ValueError("Could not find a Registro header row with a date column.")


def find_date_column_index(headers: tuple[str, ...]) -> int | None:
    for index, header in enumerate(headers):
        if normalize_header(header) in DATE_HEADERS:
            return index
    return None


def find_comment_column_index(headers: tuple[str, ...]) -> int | None:
    for index, header in enumerate(headers):
        normalized_header = normalize_header(header)
        if normalized_header in {
            "comentario",
            "comentarios",
            "comment",
            "comments",
            "nota",
            "notas",
            "observaciones",
        }:
            return index
    return None


def category_column_indexes(
    headers: tuple[str, ...], date_column_index: int | None
) -> tuple[int, ...]:
    indexes = []
    for index, header in enumerate(headers):
        normalized_header = normalize_header(header)
        if index == date_column_index:
            continue
        if not normalized_header or normalized_header in IGNORED_HEADERS:
            continue
        indexes.append(index)
    return tuple(indexes)


def infer_source_column_kind(header: str, header_cell: Any) -> str:
    if is_income_fill(header_fill_rgb(header_cell)):
        return "income"
    if is_income_header(header):
        return "income"
    return "expense"


def direction_from_source_amount(
    amount: Decimal,
    *,
    source_column_kind: str,
) -> Direction:
    if source_column_kind == "income":
        return Direction.INFLOW if amount > 0 else Direction.OUTFLOW
    return Direction.OUTFLOW if amount > 0 else Direction.INFLOW


def build_tracking_validation(
    workbook,
    *,
    tracking_sheet_name: str,
    candidates: tuple[HistoricalExcelTransactionCandidate, ...],
) -> HistoricalExcelTrackingValidation | None:
    if tracking_sheet_name not in workbook.sheetnames:
        return None

    validation_candidates = spending_tracking_validation_candidates(candidates)
    registro_totals = registro_totals_by_month_and_category(validation_candidates)
    if not registro_totals:
        return None

    sheet = workbook[tracking_sheet_name]
    rows = tuple(tuple(row) for row in sheet.iter_rows())
    candidate_years = tuple(
        sorted({candidate.transaction_date.year for candidate in validation_candidates})
    )
    tracking_result = parse_tracking_totals(
        rows,
        source_categories={
            candidate.source_category_name for candidate in validation_candidates
        },
        candidate_years=candidate_years,
    )
    if tracking_result is None:
        return None

    tracking_totals = tracking_result.totals
    registro_keys = set(registro_totals)
    tracking_keys = set(tracking_totals)
    all_keys = sorted(registro_keys | tracking_keys)
    registro_categories = {
        candidate.source_category_name for candidate in validation_candidates
    }
    seguimiento_categories = set(tracking_result.categories)
    comparisons = tuple(
        HistoricalExcelTrackingComparison(
            month_key=month_key,
            source_category_name=category_name,
            registro_amount_minor=registro_totals.get((month_key, category_name), 0),
            seguimiento_amount_minor=tracking_totals.get((month_key, category_name), 0),
        )
        for month_key, category_name in all_keys
    )
    return HistoricalExcelTrackingValidation(
        sheet_name=tracking_sheet_name,
        comparisons=comparisons,
        registro_only_categories=category_names_only_in(
            registro_categories,
            seguimiento_categories,
        ),
        seguimiento_only_categories=category_names_only_in(
            seguimiento_categories,
            registro_categories,
        ),
    )


def spending_tracking_validation_candidates(
    candidates: tuple[HistoricalExcelTransactionCandidate, ...],
) -> tuple[HistoricalExcelTransactionCandidate, ...]:
    return tuple(
        candidate
        for candidate in candidates
        if candidate.source_column_kind != "income"
    )


def registro_totals_by_month_and_category(
    candidates: tuple[HistoricalExcelTransactionCandidate, ...],
) -> dict[tuple[str, str], int]:
    totals: dict[tuple[str, str], Decimal] = defaultdict(Decimal)
    for candidate in candidates:
        key = (
            candidate.transaction_date.strftime("%Y-%m"),
            candidate.source_category_name,
        )
        totals[key] += candidate_source_amount_decimal(candidate)
    return {
        key: decimal_to_minor_units(amount)
        for key, amount in sorted(totals.items())
    }


def category_names_only_in(
    source_categories: set[str],
    other_categories: set[str],
) -> tuple[str, ...]:
    other_normalized = {
        normalize_text_for_matching(category_name)
        for category_name in other_categories
    }
    return tuple(
        sorted(
            category_name
            for category_name in source_categories
            if normalize_text_for_matching(category_name) not in other_normalized
        )
    )


def candidate_source_amount_decimal(
    candidate: HistoricalExcelTransactionCandidate,
) -> Decimal:
    if candidate.source_amount_decimal is not None:
        return candidate.source_amount_decimal
    parsed_amount = parse_excel_amount(candidate.amount_raw)
    if parsed_amount is not None:
        return parsed_amount
    return Decimal(candidate.source_signed_amount_minor) / Decimal("100")


def parse_tracking_totals(
    rows: tuple[tuple[Any, ...], ...],
    *,
    source_categories: set[str],
    candidate_years: tuple[int, ...],
) -> TrackingTotalsParseResult | None:
    row_oriented = parse_tracking_rows_by_month(
        rows,
        source_categories=source_categories,
        candidate_years=candidate_years,
    )
    if row_oriented is not None:
        return row_oriented

    column_oriented = parse_tracking_rows_by_category(
        rows,
        source_categories=source_categories,
        candidate_years=candidate_years,
    )
    return column_oriented


def parse_tracking_rows_by_month(
    rows: tuple[tuple[Any, ...], ...],
    *,
    source_categories: set[str],
    candidate_years: tuple[int, ...],
) -> TrackingTotalsParseResult | None:
    normalized_categories = normalized_category_map(source_categories)
    best_result: TrackingTotalsParseResult | None = None

    for header_row_index, header_row in enumerate(rows):
        header_values = row_values(header_row)
        tracking_categories = tracking_categories_from_month_header(header_values)
        category_columns = {
            column_index: normalized_categories.get(
                normalize_text_for_matching(str(value)),
                str(value).strip(),
            )
            for column_index, value in enumerate(header_values)
            if normalize_text_for_matching(str(value))
            in {
                normalize_text_for_matching(category_name)
                for category_name in tracking_categories
            }
        }
        if not category_columns:
            continue

        explicit_month_columns = [
            column_index
            for column_index, value in enumerate(header_values)
            if normalize_text_for_matching(str(value)) in TRACKING_MONTH_HEADERS
        ]
        month_column_candidates = explicit_month_columns or [
            column_index
            for column_index in range(max(len(header_row), 1))
            if column_index not in category_columns
        ]

        for month_column_index in month_column_candidates:
            totals = parse_tracking_month_block(
                rows,
                header_row_index=header_row_index,
                month_column_index=month_column_index,
                category_columns=category_columns,
                candidate_years=candidate_years,
            )
            if totals and explicit_month_columns:
                return TrackingTotalsParseResult(
                    totals=dict(sorted(totals.items())),
                    categories=tuple(sorted(tracking_categories)),
                )
            if best_result is None or len(totals) > len(best_result.totals):
                best_result = TrackingTotalsParseResult(
                    totals=dict(totals),
                    categories=tuple(sorted(tracking_categories)),
                )

    if best_result is None:
        return None
    return TrackingTotalsParseResult(
        totals=dict(sorted(best_result.totals.items())),
        categories=best_result.categories,
    )


def parse_tracking_month_block(
    rows: tuple[tuple[Any, ...], ...],
    *,
    header_row_index: int,
    month_column_index: int,
    category_columns: dict[int, str],
    candidate_years: tuple[int, ...],
) -> dict[tuple[str, str], int]:
    totals: dict[tuple[str, str], int] = defaultdict(int)
    started_month_rows = False

    for row in rows[header_row_index + 1 :]:
        month_key = parse_tracking_month(
            cell_value(row, month_column_index),
            candidate_years=candidate_years,
        )
        if month_key is None:
            if started_month_rows:
                break
            if row_has_any_value(row):
                continue
            continue

        started_month_rows = True
        for category_column_index, category_name in category_columns.items():
            amount = parse_excel_amount(cell_value(row, category_column_index))
            if amount is None or amount == 0:
                continue
            totals[(month_key, category_name)] += decimal_to_minor_units(amount)

    return dict(totals)


def parse_tracking_rows_by_category(
    rows: tuple[tuple[Any, ...], ...],
    *,
    source_categories: set[str],
    candidate_years: tuple[int, ...],
) -> TrackingTotalsParseResult | None:
    normalized_categories = normalized_category_map(source_categories)
    best_result: TrackingTotalsParseResult | None = None

    for header_row_index, header_row in enumerate(rows):
        month_columns = {
            column_index: month_key
            for column_index, value in enumerate(row_values(header_row))
            if (
                month_key := parse_tracking_month(
                    value,
                    candidate_years=candidate_years,
                )
            )
            is not None
        }
        if not month_columns:
            continue

        for category_column_index in range(max(len(header_row), 1)):
            if category_column_index in month_columns:
                continue
            totals: dict[tuple[str, str], int] = defaultdict(int)
            tracking_categories: set[str] = set()
            for row in rows[header_row_index + 1 :]:
                category_value = cell_value(row, category_column_index)
                normalized_category = normalize_text_for_matching(str(category_value))
                if normalized_category not in normalized_categories:
                    continue
                category_name = normalized_categories[normalized_category]
                tracking_categories.add(category_name)
                for month_column_index, month_key in month_columns.items():
                    amount = parse_excel_amount(cell_value(row, month_column_index))
                    if amount is None or amount == 0:
                        continue
                    totals[(month_key, category_name)] += decimal_to_minor_units(amount)
            if best_result is None or len(totals) > len(best_result.totals):
                best_result = TrackingTotalsParseResult(
                    totals=dict(totals),
                    categories=tuple(sorted(tracking_categories)),
                )

    if best_result is None:
        return None
    return TrackingTotalsParseResult(
        totals=dict(sorted(best_result.totals.items())),
        categories=best_result.categories,
    )


def tracking_categories_from_month_header(header_values: tuple[Any, ...]) -> tuple[str, ...]:
    categories = []
    for value in header_values:
        if value is None:
            continue
        label = str(value).strip()
        normalized_label = normalize_text_for_matching(label)
        if not normalized_label:
            continue
        if normalized_label in TRACKING_MONTH_HEADERS:
            continue
        if normalized_label in IGNORED_HEADERS:
            continue
        categories.append(label)
    return tuple(categories)


def normalized_category_map(source_categories: set[str]) -> dict[str, str]:
    return {
        normalize_text_for_matching(category_name): category_name
        for category_name in source_categories
    }


def row_values(row: tuple[Any, ...]) -> tuple[Any, ...]:
    return tuple(cell_value(row, index) for index in range(len(row)))


def row_has_any_value(row: tuple[Any, ...]) -> bool:
    return any(cell_value(row, index) not in (None, "") for index in range(len(row)))


def parse_tracking_month(
    value: Any,
    *,
    candidate_years: tuple[int, ...],
) -> str | None:
    parsed_date = parse_excel_date(value)
    if parsed_date is not None:
        return parsed_date.strftime("%Y-%m")

    if value is None:
        return None
    text = normalize_text_for_matching(str(value))
    if not text:
        return None

    for date_format in ("%Y-%m", "%m/%Y", "%m-%Y"):
        try:
            parsed = datetime.strptime(text, date_format)
            return parsed.strftime("%Y-%m")
        except ValueError:
            continue

    if text in TRACKING_MONTH_HEADERS or text in TRACKING_CATEGORY_HEADERS:
        return None

    if text in MONTH_NAMES and len(candidate_years) == 1:
        return f"{candidate_years[0]:04d}-{MONTH_NAMES[text]:02d}"

    parts = text.replace("/", " ").replace("-", " ").split()
    month = next((MONTH_NAMES[part] for part in parts if part in MONTH_NAMES), None)
    year = next((int(part) for part in parts if part.isdigit() and len(part) == 4), None)
    if month is not None and year is not None:
        return f"{year:04d}-{month:02d}"
    if month is not None and len(candidate_years) == 1:
        return f"{candidate_years[0]:04d}-{month:02d}"

    return None


def is_income_header(header: str) -> bool:
    normalized_header = normalize_text_for_matching(header)
    income_keywords = {
        "dividendo",
        "dividendos",
        "income",
        "ingreso",
        "ingresos",
        "interes",
        "intereses",
        "nomina",
        "payroll",
        "salario",
        "salary",
    }
    return any(keyword in normalized_header for keyword in income_keywords)


def header_fill_rgb(cell: Any) -> str | None:
    fill = getattr(cell, "fill", None)
    if fill is None or fill.fill_type is None:
        return None
    color = fill.fgColor
    if color is None or color.type != "rgb" or color.rgb is None:
        return None
    rgb = str(color.rgb)
    return rgb[-6:].upper()


def is_income_fill(rgb: str | None) -> bool:
    if rgb is None or len(rgb) != 6:
        return False
    try:
        red = int(rgb[0:2], 16)
        green = int(rgb[2:4], 16)
        blue = int(rgb[4:6], 16)
    except ValueError:
        return False
    return green > red and green > blue


def parse_excel_date(value: Any) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    for date_format in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue
    return None


def parse_excel_amount(value: Any) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, int | float | Decimal):
        return Decimal(str(value))
    text = str(value).strip()
    if not text:
        return None
    normalized = normalize_decimal_text(text)
    try:
        return Decimal(normalized)
    except InvalidOperation:
        return None


def normalize_decimal_text(value: str) -> str:
    text = value.replace("€", "").replace(" ", "")
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            return text.replace(".", "").replace(",", ".")
        return text.replace(",", "")
    if "," in text:
        return text.replace(",", ".")
    return text


def decimal_to_minor_units(amount: Decimal) -> int:
    return int((amount * Decimal("100")).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def normalize_header(value: str) -> str:
    return value.strip().lower()


def normalize_text_for_matching(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value.strip().lower())
    return "".join(character for character in normalized if not unicodedata.combining(character))


def cell_value(row: tuple[Any, ...], index: int | None) -> Any:
    if index is None or index >= len(row):
        return None
    cell = row[index]
    return cell.value if hasattr(cell, "value") else cell


def string_or_none(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def payload_from_row(headers: tuple[str, ...], row: tuple[Any, ...]) -> dict[str, str | None]:
    payload = {}
    for index, header in enumerate(headers):
        if not header:
            continue
        payload[header] = string_or_none(cell_value(row, index))
    return payload


__all__ = [
    "HistoricalExcelDryRunImporter",
    "HistoricalExcelPreview",
    "HistoricalExcelTransactionCandidate",
]
