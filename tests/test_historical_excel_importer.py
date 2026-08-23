from datetime import date

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from lxcell.enums.core_enums import Direction
from lxcell.importers import HistoricalExcelDryRunImporter


def test_historical_excel_preview_reads_registro_candidates(tmp_path):
    workbook_path = tmp_path / "historical_sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Registro"
    sheet.append(["Sample historical workbook"])
    sheet.append([])
    sheet.append(["Fecha", "Supermercado", "Nómina", "Total", "Comentarios"])
    sheet["C3"].fill = PatternFill("solid", fgColor="92D050")
    sheet.append([date(2026, 1, 1), 12.34, None, 12.34, "Compra semanal"])
    sheet.append([date(2026, 1, 2), None, 1000, 1000, "Ingreso mensual"])
    sheet.append([date(2026, 1, 3), -2.5, None, -2.5, "Devolución"])
    sheet.append([date(2026, 1, 4), 0, None, 0, "Sin movimiento"])
    tracking_sheet = workbook.create_sheet("Seguimiento")
    tracking_sheet.append(["Mes", "Supermercado"])
    tracking_sheet.append([date(2026, 1, 1), 9.84])
    workbook.save(workbook_path)

    preview = HistoricalExcelDryRunImporter().preview(workbook_path)

    assert preview.source_file_name == "historical_sample.xlsx"
    assert len(preview.source_file_hash) == 64
    assert preview.sheet_name == "Registro"
    assert preview.header_row_number == 3
    assert preview.date_column_name == "Fecha"
    assert preview.transaction_count == 3
    assert preview.source_categories == ("Nómina", "Supermercado")
    assert preview.totals_by_category_minor == {
        "Nómina": 100000,
        "Supermercado": -984,
    }
    assert preview.totals_by_month_minor == {"2026-01": 99016}
    assert preview.source_totals_by_category_minor == {
        "Nómina": 100000,
        "Supermercado": 984,
    }
    assert preview.source_totals_by_month_minor == {"2026-01": 100984}
    assert preview.tracking_validation is not None
    assert preview.tracking_validation.ok_count == 1
    assert preview.tracking_validation.difference_count == 0
    assert preview.tracking_validation.registro_only_categories == ()
    assert preview.tracking_validation.seguimiento_only_categories == ()
    assert preview.ignored_row_numbers == (7,)

    expense, income, refund = preview.candidates
    assert expense.transaction_date == date(2026, 1, 1)
    assert expense.source_category_name == "Supermercado"
    assert expense.amount_minor == 1234
    assert expense.source_amount_minor == 1234
    assert expense.source_column_kind == "expense"
    assert expense.direction == Direction.OUTFLOW
    assert expense.description_raw == "Compra semanal"
    assert expense.payload_raw["Comentarios"] == "Compra semanal"

    assert income.transaction_date == date(2026, 1, 2)
    assert income.source_category_name == "Nómina"
    assert income.amount_minor == 100000
    assert income.source_amount_minor == 100000
    assert income.source_column_kind == "income"
    assert income.direction == Direction.INFLOW

    assert refund.transaction_date == date(2026, 1, 3)
    assert refund.source_category_name == "Supermercado"
    assert refund.amount_minor == 250
    assert refund.source_amount_minor == -250
    assert refund.source_column_kind == "expense"
    assert refund.direction == Direction.INFLOW


def test_historical_excel_preview_reads_category_oriented_seguimiento(tmp_path):
    workbook_path = tmp_path / "historical_sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Registro"
    sheet.append(["Fecha", "Category A"])
    sheet.append([date(2026, 1, 1), 10])
    tracking_sheet = workbook.create_sheet("Seguimiento")
    tracking_sheet.append(["Categoría", date(2026, 1, 1)])
    tracking_sheet.append(["Category A", 10.01])
    workbook.save(workbook_path)

    preview = HistoricalExcelDryRunImporter().preview(workbook_path)

    assert preview.tracking_validation is not None
    comparison = preview.tracking_validation.comparisons[0]
    assert comparison.month_key == "2026-01"
    assert comparison.source_category_name == "Category A"
    assert comparison.registro_amount_minor == 1000
    assert comparison.seguimiento_amount_minor == 1001
    assert comparison.difference_minor == -1
    assert comparison.status == "ok"


def test_historical_excel_preview_uses_only_first_seguimiento_monthly_table(tmp_path):
    workbook_path = tmp_path / "historical_sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Registro"
    sheet.append(["Fecha", "Casa", "Ocio"])
    sheet.append([date(2025, 1, 1), 2428.54, 430.52])

    tracking_sheet = workbook.create_sheet("Seguimiento")
    tracking_sheet.append(["Read-only warning"])
    tracking_sheet.append([])
    tracking_sheet.append(["Mes", "Casa", "Ocio"])
    tracking_sheet.append([date(2025, 1, 1), 2428.54, 430.52])
    tracking_sheet.append(["Año completo", 20984.28, 3311.07])
    tracking_sheet.append(["Superávit/Déficit acumulado (desglose)", -1924.28, 48.93])
    tracking_sheet.append([])
    tracking_sheet.append(["Mes", "Salario", "Total Gasto"])
    tracking_sheet.append([date(2025, 1, 1), 5284.9, 4378.79])
    tracking_sheet.append([])
    tracking_sheet.append(["Superávit/Déficit mensual desglosado"])
    tracking_sheet.append(["Mes", "Casa", "Ocio"])
    tracking_sheet.append([date(2025, 1, 1), -1098.54, 397.22])
    workbook.save(workbook_path)

    preview = HistoricalExcelDryRunImporter().preview(workbook_path)

    assert preview.tracking_validation is not None
    comparisons = {
        (comparison.month_key, comparison.source_category_name): comparison
        for comparison in preview.tracking_validation.comparisons
    }
    assert comparisons[("2025-01", "Casa")].seguimiento_amount_minor == 242854
    assert comparisons[("2025-01", "Casa")].status == "ok"
    assert comparisons[("2025-01", "Ocio")].seguimiento_amount_minor == 43052
    assert comparisons[("2025-01", "Ocio")].status == "ok"


def test_historical_excel_validation_rounds_registro_after_aggregation(tmp_path):
    workbook_path = tmp_path / "historical_sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Registro"
    sheet.append(["Fecha", "Supermercado"])
    for day in range(1, 7):
        sheet.append([date(2025, 1, day), 0.015])

    tracking_sheet = workbook.create_sheet("Seguimiento")
    tracking_sheet.append(["Mes", "Supermercado"])
    tracking_sheet.append([date(2025, 1, 1), 0.09])
    workbook.save(workbook_path)

    preview = HistoricalExcelDryRunImporter().preview(workbook_path)

    assert preview.tracking_validation is not None
    comparison = preview.tracking_validation.comparisons[0]
    assert comparison.registro_amount_minor == 9
    assert comparison.seguimiento_amount_minor == 9
    assert comparison.status == "ok"


def test_historical_excel_validation_ignores_zero_tracking_cells(tmp_path):
    workbook_path = tmp_path / "historical_sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Registro"
    sheet.append(["Fecha", "Supermercado"])
    sheet.append([date(2025, 1, 1), 10])

    tracking_sheet = workbook.create_sheet("Seguimiento")
    tracking_sheet.append(["Mes", "Supermercado"])
    tracking_sheet.append([date(2025, 1, 1), 10])
    tracking_sheet.append([date(2025, 2, 1), 0])
    workbook.save(workbook_path)

    preview = HistoricalExcelDryRunImporter().preview(workbook_path)

    assert preview.tracking_validation is not None
    assert [
        (comparison.month_key, comparison.source_category_name)
        for comparison in preview.tracking_validation.comparisons
    ] == [("2025-01", "Supermercado")]
    assert preview.tracking_validation.seguimiento_only_categories == ()


def test_historical_excel_validation_excludes_income_categories_from_first_tracking_table(
    tmp_path,
):
    workbook_path = tmp_path / "historical_sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Registro"
    sheet.append(["Fecha", "Supermercado", "Salary"])
    sheet["C1"].fill = PatternFill("solid", fgColor="92D050")
    sheet.append([date(2025, 1, 1), 25, 1000])

    tracking_sheet = workbook.create_sheet("Seguimiento")
    tracking_sheet.append(["Mes", "Supermercado"])
    tracking_sheet.append([date(2025, 1, 1), 25])
    workbook.save(workbook_path)

    preview = HistoricalExcelDryRunImporter().preview(workbook_path)

    assert preview.source_categories == ("Salary", "Supermercado")
    assert preview.tracking_validation is not None
    assert [
        (comparison.month_key, comparison.source_category_name)
        for comparison in preview.tracking_validation.comparisons
    ] == [("2025-01", "Supermercado")]
    assert preview.tracking_validation.registro_only_categories == ()
    assert preview.tracking_validation.seguimiento_only_categories == ()


def test_historical_excel_preview_rejects_missing_registro_sheet(tmp_path):
    workbook_path = tmp_path / "historical_sample.xlsx"
    workbook = Workbook()
    workbook.active.title = "Other"
    workbook.save(workbook_path)

    try:
        HistoricalExcelDryRunImporter().preview(workbook_path)
    except ValueError as exc:
        assert "Registro" in str(exc)
    else:
        raise AssertionError("Expected missing Registro sheet to fail.")
